import streamlit as st
import requests
import re

from pathlib import Path
from io import BytesIO

from docx import Document
from docx.shared import Inches, Pt
from docx.enum.table import WD_TABLE_ALIGNMENT

from fpdf import FPDF


# ============================================================
# PAGE CONFIGURATION
# ============================================================

st.set_page_config(
    page_title="LegalEase",
    page_icon="⚖️",
    layout="wide"
)


# ============================================================
# PROJECT PATHS
# ============================================================

BASE_DIR = Path(__file__).resolve().parents[1]

LOGO_PATH = BASE_DIR / "image" / "logo.png"

API_URL = "http://127.0.0.1:8000/generate"


# ============================================================
# CLEAN GENERATED DOCUMENT
# ============================================================

def clean_document(text):

    if not text:
        return ""

    # Remove Markdown headings
    text = re.sub(
        r'^\s*#+\s*',
        '',
        text,
        flags=re.MULTILINE
    )

    # Remove bold / italic Markdown
    text = text.replace("**", "")
    text = text.replace("*", "")

    # Remove Markdown horizontal lines
    text = re.sub(
        r'^\s*-{3,}\s*$',
        '',
        text,
        flags=re.MULTILINE
    )

    # Remove excessive blank lines
    text = re.sub(
        r'\n{3,}',
        '\n\n',
        text
    )

    return text.strip()


# ============================================================
# UNICODE CLEANING FOR PDF
# ============================================================

def clean_pdf_text(text):

    if not text:
        return ""

    replacements = {
        "’": "'",
        "‘": "'",
        "“": '"',
        "”": '"',
        "–": "-",
        "—": "-",
        "…": "...",
        "•": "-",
        "\u00a0": " ",
        "\u2011": "-",
        "\u2012": "-",
        "\u2013": "-",
        "\u2014": "-",
        "\u2212": "-",
        "\u2022": "-",
        "\u00ad": "",
    }

    for old, new in replacements.items():
        text = text.replace(old, new)

    # Convert remaining unsupported characters
    text = text.encode(
        "latin-1",
        errors="replace"
    ).decode("latin-1")

    return text


# ============================================================
# CREATE DOCX
# ============================================================

def create_docx(document_text, terms_text=""):

    doc = Document()

    # --------------------------------------------------------
    # Logo
    # --------------------------------------------------------

    if LOGO_PATH.exists():

        try:

            paragraph = doc.add_paragraph()

            paragraph.alignment = 1

            run = paragraph.add_run()

            run.add_picture(
                str(LOGO_PATH),
                width=Inches(2)
            )

        except Exception as e:

            print(
                "Logo could not be added:",
                e
            )

    # --------------------------------------------------------
    # Document Title
    # --------------------------------------------------------

    title = doc.add_paragraph()

    title.alignment = 1

    run = title.add_run(
        "LEGAL DOCUMENT"
    )

    run.bold = True
    run.font.name = "Times New Roman"
    run.font.size = Pt(16)

    # --------------------------------------------------------
    # Generated Document Content
    # --------------------------------------------------------

    for line in document_text.split("\n"):

        line = line.strip()

        if not line:
            continue

        paragraph = doc.add_paragraph()

        run = paragraph.add_run(line)

        run.font.name = "Times New Roman"
        run.font.size = Pt(12)

    # --------------------------------------------------------
    # Terms & Conditions Table
    # --------------------------------------------------------

    if terms_text.strip():

        doc.add_paragraph()

        heading = doc.add_paragraph()

        run = heading.add_run(
            "Terms & Conditions"
        )

        run.bold = True
        run.font.name = "Times New Roman"
        run.font.size = Pt(13)

        # Split terms by semicolon
        terms = [
            term.strip()
            for term in terms_text.split(";")
            if term.strip()
        ]

        if terms:

            table = doc.add_table(
                rows=1,
                cols=2
            )

            table.style = "Table Grid"

            table.alignment = (
                WD_TABLE_ALIGNMENT.CENTER
            )

            # Header
            header_cells = table.rows[0].cells

            header_cells[0].text = "S.No."
            header_cells[1].text = "Terms & Conditions"

            for cell in header_cells:

                for paragraph in cell.paragraphs:

                    for run in paragraph.runs:

                        run.bold = True
                        run.font.name = "Times New Roman"
                        run.font.size = Pt(11)

            # Table rows
            for index, term in enumerate(
                terms,
                start=1
            ):

                cells = table.add_row().cells

                cells[0].text = str(index)
                cells[1].text = term

                for cell in cells:

                    for paragraph in cell.paragraphs:

                        for run in paragraph.runs:

                            run.font.name = "Times New Roman"
                            run.font.size = Pt(11)

    # --------------------------------------------------------
    # Footer
    # --------------------------------------------------------

    section = doc.sections[0]

    footer = section.footer

    footer_paragraph = footer.paragraphs[0]

    footer_paragraph.alignment = 1

    footer_paragraph.text = (
        "LegalEase - AI-Powered Legal Document Generator"
    )

    for run in footer_paragraph.runs:

        run.font.name = "Times New Roman"
        run.font.size = Pt(9)

    # --------------------------------------------------------
    # Create DOCX
    # --------------------------------------------------------

    docx_file = BytesIO()

    doc.save(docx_file)

    docx_file.seek(0)

    return docx_file


# ============================================================
# CREATE PDF
# ============================================================

def create_pdf(
    document_text,
    terms_text=""
):

    class LegalEasePDF(FPDF):

        # ----------------------------------------------------
        # Header
        # ----------------------------------------------------

        def header(self):

            if LOGO_PATH.exists():

                try:

                    self.image(
                        str(LOGO_PATH),
                        x=10,
                        y=8,
                        w=35
                    )

                except Exception as e:

                    print(
                        "PDF logo error:",
                        e
                    )

            self.set_font(
                "Arial",
                "B",
                14
            )

            self.cell(
                w=0,
                h=10,
                text="LegalEase",
                align="R"
            )

            self.ln(18)

        # ----------------------------------------------------
        # Footer
        # ----------------------------------------------------

        def footer(self):

            self.set_y(-15)

            self.set_font(
                "Arial",
                size=8
            )

            self.cell(
                w=0,
                h=10,
                text=(
                    "LegalEase - AI-Powered Legal "
                    "Document Generator"
                ),
                align="C"
            )

    # --------------------------------------------------------
    # Create PDF
    # --------------------------------------------------------

    pdf = LegalEasePDF()

    pdf.set_auto_page_break(
        auto=True,
        margin=20
    )

    pdf.set_margins(
        left=15,
        top=15,
        right=15
    )

    pdf.add_page()

    # Clean Unicode
    document_text = clean_pdf_text(
        document_text
    )

    # --------------------------------------------------------
    # Document Content
    # --------------------------------------------------------

    for line in document_text.split("\n"):

        line = line.strip()

        if not line:

            pdf.ln(4)

            continue

        # ----------------------------------------------------
        # Heading
        # ----------------------------------------------------

        if (
            line.isupper()
            or re.match(
                r"^\d+\.",
                line
            )
        ):

            pdf.set_font(
                "Arial",
                "B",
                12
            )

            pdf.multi_cell(
                w=pdf.epw,
                h=8,
                text=line,
                border=0,
                new_x="LMARGIN",
                new_y="NEXT",
                wrapmode="CHAR"
            )

            pdf.ln(2)

        # ----------------------------------------------------
        # Bullet
        # ----------------------------------------------------

        elif line.startswith("-"):

            pdf.set_font(
                "Arial",
                size=11
            )

            pdf.multi_cell(
                w=pdf.epw,
                h=7,
                text=line,
                border=0,
                new_x="LMARGIN",
                new_y="NEXT",
                wrapmode="CHAR"
            )

        # ----------------------------------------------------
        # Normal paragraph
        # ----------------------------------------------------

        else:

            pdf.set_font(
                "Arial",
                size=11
            )

            pdf.multi_cell(
                w=pdf.epw,
                h=7,
                text=line,
                border=0,
                new_x="LMARGIN",
                new_y="NEXT",
                wrapmode="CHAR"
            )

    # --------------------------------------------------------
    # Terms & Conditions
    # --------------------------------------------------------

    if terms_text.strip():

        terms = [
            term.strip()
            for term in terms_text.split(";")
            if term.strip()
        ]

        if terms:

            pdf.ln(5)

            pdf.set_font(
                "Arial",
                "B",
                12
            )

            pdf.multi_cell(
                w=pdf.epw,
                h=8,
                text="TERMS & CONDITIONS",
                border=0,
                new_x="LMARGIN",
                new_y="NEXT",
                wrapmode="CHAR"
            )

            pdf.ln(2)

            for index, term in enumerate(
                terms,
                start=1
            ):

                term = clean_pdf_text(term)

                bullet_text = (
                    f"{index}. {term}"
                )

                pdf.set_font(
                    "Arial",
                    size=11
                )

                pdf.multi_cell(
                    w=pdf.epw,
                    h=7,
                    text=bullet_text,
                    border=0,
                    new_x="LMARGIN",
                    new_y="NEXT",
                    wrapmode="CHAR"
                )

    # --------------------------------------------------------
    # Create PDF Bytes
    # --------------------------------------------------------

    pdf_bytes = pdf.output()

    pdf_file = BytesIO(
        bytes(pdf_bytes)
    )

    pdf_file.seek(0)

    return pdf_file


# ============================================================
# SESSION STATE
# ============================================================

if "generated_document" not in st.session_state:

    st.session_state.generated_document = ""


# ============================================================
# PAGE TITLE
# ============================================================

st.title(
    "⚖️ LegalEase"
)

st.subheader(
    "AI-Powered Legal Document Generator"
)

st.write(
    "Generate customizable legal documents "
    "using AI and download them in TXT, DOCX, "
    "or PDF format."
)


# ============================================================
# INPUT SECTION
# ============================================================

st.markdown(
    "### Document Details"
)


# ------------------------------------------------------------
# Document Type
# ------------------------------------------------------------

document_type = st.selectbox(
    "Document Type",
    [
        "Employment Contract",
        "Lease Agreement",
        "Non-Disclosure Agreement (NDA)",
        "Service Agreement"
    ]
)


# ------------------------------------------------------------
# Parties
# ------------------------------------------------------------

parties = st.text_area(
    "Parties",
    placeholder=(
        "Example: ABC Company and John Doe"
    ),
    height=100
)


# ------------------------------------------------------------
# Terms
# ------------------------------------------------------------

terms = st.text_area(
    "Terms and Conditions",
    placeholder=(
        "Example: Salary is 30000; "
        "Working hours are 9 AM to 6 PM; "
        "Confidentiality must be maintained"
    ),
    height=120
)


# ------------------------------------------------------------
# Dates
# ------------------------------------------------------------

dates = st.text_input(
    "Important Dates",
    placeholder="Example: 10/10/2026"
)


# ============================================================
# GENERATE BUTTON
# ============================================================

if st.button(
    "Generate Document",
    type="primary",
    use_container_width=False
):

    # Clear previous document before generating a new one.
    st.session_state.generated_document = ""

    # Basic validation
    if not parties.strip():

        st.warning(
            "Please enter the parties."
        )

    elif not terms.strip():

        st.warning(
            "Please enter the terms and conditions."
        )

    elif not dates.strip():

        st.warning(
            "Please enter the important dates."
        )

    else:

        try:

            with st.spinner(
                "⏳ Generating document... Please wait."
            ):

                response = requests.post(
                    API_URL,
                    json={
                        "document_type": document_type,
                        "parties": parties,
                        "terms": terms,
                        "dates": dates
                    },
                    timeout=120
                )

            # ==================================================
            # SUCCESSFUL HTTP RESPONSE
            # ==================================================

            if response.status_code == 200:

                response_data = response.json()

                # ------------------------------------------------
                # IMPORTANT:
                # Check backend success first
                # ------------------------------------------------

                if response_data.get("success") is False:

                    error_message = response_data.get(
                        "error",
                        "Unknown error occurred."
                    )

                    st.error(
                        "Document generation failed:"
                    )

                    st.code(
                        error_message
                    )

                else:

                    generated_document = (
                        response_data.get(
                            "document",
                            ""
                        )
                    )

                    # Check actual document content
                    if (
                        generated_document
                        and generated_document.strip()
                    ):

                        generated_document = (
                            clean_document(
                                generated_document
                            )
                        )

                        st.session_state.generated_document = (
                            generated_document
                        )

                        st.success(
                            "Document generated successfully!"
                        )

                    else:

                        st.error(
                            "The API returned an empty document."
                        )

            # ==================================================
            # NON-200 RESPONSE
            # ==================================================

            else:

                st.error(
                    f"Error generating document: "
                    f"{response.status_code}"
                )

                # Show backend response
                try:

                    response_data = response.json()

                    if response_data.get("error"):

                        st.code(
                            response_data["error"]
                        )

                    elif response_data.get("detail"):

                        st.code(
                            str(
                                response_data["detail"]
                            )
                        )

                    else:

                        st.code(
                            response.text
                        )

                except Exception:

                    st.code(
                        response.text
                    )

        except requests.exceptions.ConnectionError:

            st.error(
                "Cannot connect to FastAPI backend. "
                "Please make sure the FastAPI server is running."
            )

        except requests.exceptions.Timeout:

            st.error(
                "The request timed out. "
                "Please try again."
            )

        except Exception as e:

            st.error(
                f"Unexpected error: {e}"
            )


# ============================================================
# GENERATED DOCUMENT
# ============================================================

if st.session_state.generated_document:

    st.markdown(
        "### Generated Document"
    )

    # ========================================================
    # EDITABLE DOCUMENT
    # ========================================================

    edited_document = st.text_area(
        "Edit your document if required:",
        value=st.session_state.generated_document,
        height=600
    )

    # Keep edited version
    st.session_state.generated_document = edited_document

    # ========================================================
    # DOWNLOAD SECTION
    # ========================================================

    st.markdown(
        "### Download"
    )

    # ========================================================
    # CREATE DOCX
    # ========================================================

    try:

        docx_file = create_docx(
            st.session_state.generated_document,
            terms
        )

        docx_data = docx_file.getvalue()

    except Exception as e:

        docx_data = None

        st.error(
            f"DOCX creation error: {e}"
        )

    # ========================================================
    # CREATE PDF
    # ========================================================

    try:

        pdf_file = create_pdf(
            st.session_state.generated_document,
            terms
        )

        pdf_data = pdf_file.getvalue()

    except Exception as e:

        pdf_data = None

        st.error(
            f"PDF creation error: {e}"
        )

    # ========================================================
    # DOWNLOAD BUTTONS
    # ========================================================

    col1, col2, col3 = st.columns(3)

    # --------------------------------------------------------
    # TXT
    # --------------------------------------------------------

    with col1:

       st.download_button(
    label="📄 Download TXT",
    data=(
        st.session_state.generated_document
        .encode("utf-8")
    ),
    file_name="LegalEase_Document.txt",
    mime="text/plain",
    use_container_width=True,
    key="download_txt",
    on_click="ignore"
)

    # --------------------------------------------------------
    # DOCX
    # --------------------------------------------------------

    with col2:

        if docx_data:

           st.download_button(
    label="📝 Download DOCX",
    data=docx_data,
    file_name="LegalEase_Document.docx",
    mime=(
        "application/vnd.openxmlformats-"
        "officedocument.wordprocessingml.document"
    ),
    use_container_width=True,
    key="download_docx",
    on_click="ignore"
)

    # --------------------------------------------------------
    # PDF
    # --------------------------------------------------------

    with col3:

        if pdf_data:

          st.download_button(
    label="📕 Download PDF",
    data=pdf_data,
    file_name="LegalEase_Document.pdf",
    mime="application/pdf",
    use_container_width=True,
    key="download_pdf",
    on_click="ignore"
)